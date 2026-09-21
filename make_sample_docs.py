"""Create the optional PDF + DOCX sample documents in knowledge_base/.   python scripts/make_sample_docs.py
Then run: python -m src.ingest      (needs fpdf2 + python-docx; both are in requirements-dev.txt)"""
from pathlib import Path

KB = Path(__file__).resolve().parent.parent / "knowledge_base"

GIFT = [
    ("Gift card balance", "ShopEase gift card balances never expire. A gift card is applied at checkout and can be combined with another payment method. Gift cards cannot be returned or exchanged for cash."),
    ("Lost or stolen gift cards", "Report a lost or stolen gift card within 7 days. Only the unused balance can be recovered and it is moved to a new card."),
    ("Promotional codes", "Each order accepts one promotional code. Codes cannot be combined with other codes and expire on the date shown in the offer."),
]
SELLERS = [
    ("Items sold by third-party sellers", "Some items are sold and shipped by independent sellers. The seller name is shown on the product page and in Your Orders. Seller items follow the ShopEase return window unless the product page states a longer one."),
    ("ShopEase Guarantee", "If a seller has not resolved a problem within 7 days of your report, contact support. We will review the case and may refund the order."),
]


def make_pdf(path, title, sections):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
    for head, body in sections:
        pdf.set_font("Helvetica", "B", 12)
        pdf.multi_cell(0, 8, head, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 6, body, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)
    pdf.output(str(path))


def make_docx(path, title, sections):
    from docx import Document

    d = Document()
    d.add_heading(title, level=1)
    for head, body in sections:
        d.add_heading(head, level=2)
        d.add_paragraph(body)
    d.save(str(path))


if __name__ == "__main__":
    make_pdf(KB / "gift_cards_and_promotions.pdf", "Gift Cards and Promotions", GIFT)
    make_docx(KB / "marketplace_sellers.docx", "Marketplace Sellers", SELLERS)
    print("Created knowledge_base/gift_cards_and_promotions.pdf and marketplace_sellers.docx; now run: python -m src.ingest")
